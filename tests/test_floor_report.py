"""Real saved-result reports; faults and optional CAD are explicitly simulated."""

from copy import deepcopy
import hashlib
import json
from pathlib import Path
import subprocess
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from docx import Document
from agent.controller import AgentController
from agent.state import AgentState
from core import ToolRegistry, ToolResult
from examples.floor_report import ReportRequestParser
from tests.test_agent_controller import SimulatedCAD
from tools.floor.check_adapter import FloorCheckError
from tools.floor.design_adapter import canonical_hash
from tools.floor.plugin import register_floor_report_workflow
from tools.floor.report_adapter import FloorReportError

ROOT = Path(__file__).resolve().parents[1]
CONTEXT = dict(unit_system='SI',design_code='GB')


class FloorReportTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)/'中文 计算书'
        self.registry=ToolRegistry();self.backend=SimulatedCAD()
        self.workflow=register_floor_report_workflow(self.registry,self.root,cad_backend=self.backend)
        self.tool=self.registry.get('generate_floor_report')
        self.model=json.loads((ROOT/'legacy/rc_floor/demo_a.json').read_text(encoding='utf-8-sig'))
        self.parameters=dict(input_mode='explicit',model=self.model)
        self.design=self.design_case(self.parameters)
        self.ref=self.design.metadata['design_result_ref']

    def design_case(self,parameters):
        result=self.registry.get('design_floor_system').execute(dict(project_id='qa',tool='design_floor_system',context=CONTEXT,parameters=parameters))
        self.assertTrue(result.success,result.errors);return result

    def report(self,ref=None,project='qa',**extra):
        return self.tool.execute(dict(project_id=project,tool=self.tool.name,context=CONTEXT,
            parameters=dict(design_result_ref=ref or self.ref,**extra)))

    def test_three_original_cases_generate_from_saved_results_and_keep_sources(self):
        for name in ('demo_a','sample1','changed'):
            model=json.loads((ROOT/'legacy/rc_floor'/f'{name}.json').read_text(encoding='utf-8-sig'))
            design=self.design_case(dict(input_mode='explicit',model=model));ref=design.metadata['design_result_ref']
            snapshot=self.tool.store.root/(ref+'.json');before=snapshot.read_bytes()
            with patch.object(self.registry.get('design_floor_system'),'execute',side_effect=AssertionError('No redesign')):
                result=self.report(ref)
            self.assertTrue(result.success,result.errors)
            self.assertEqual(snapshot.read_bytes(),before)
            self.assertEqual(result.result['design_sha256'],design.metadata['legacy_result_sha256'])
            path=Path(next(a['path'] for a in result.artifacts if a['type']=='docx'))
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),result.result['docx_sha256'])
            doc=Document(path);text='\n'.join(p.text for p in doc.paragraphs)
            self.assertEqual(doc.styles['toc 1'].style_id,'TOC1')
            self.assertEqual(doc.styles['toc 2'].font.size.pt,9.5)
            self.assertIn(ref,text);self.assertIn('附录 独立截面校核记录',text)
            for chapter in design.result['legacy_result']['chapters']:self.assertIn(chapter['title'],text)
            for item in design.result['legacy_result']['chapters'][0]['items']:
                if item['type']=='text':self.assertIn(item['text'],text)
            self.assertEqual(len(doc.inline_shapes),result.result['images'])
            self.assertIsNone(result.result['revision_id'])

    def test_repeat_reports_do_not_overwrite_completed_book_or_evidence(self):
        first=self.report();self.assertTrue(first.success,first.errors)
        hashes={Path(a['path']):hashlib.sha256(Path(a['path']).read_bytes()).hexdigest() for a in first.artifacts}
        second=self.report();self.assertTrue(second.success,second.errors)
        self.assertNotEqual(first.result['report_id'],second.result['report_id'])
        self.assertTrue(all(hashlib.sha256(p.read_bytes()).hexdigest()==h for p,h in hashes.items()))

    def test_bad_reference_project_checksum_and_extra_paths_never_start_report(self):
        with patch.object(self.tool.adapter,'generate') as generate:
            for ref,project in (('../escape','qa'),(self.ref,'other'),('floor-'+'0'*32,'qa')):
                self.assertFalse(self.report(ref,project).success)
            self.assertFalse(self.report(path='arbitrary.docx').success)
            self.assertFalse(self.report(revision_id='../bad').success)
            snapshot=self.tool.store.root/(self.ref+'.json')
            raw=json.loads(snapshot.read_text(encoding='utf-8'));raw['checksum']='0'*64
            snapshot.write_text(json.dumps(raw),encoding='utf-8')
            self.assertEqual(self.report().errors[0]['code'],'invalid_design_reference')
            generate.assert_not_called()

    def test_real_under_reinforced_snapshot_cannot_generate_a_pass_book(self):
        data=self.tool.store.load(self.ref,project_id='qa').to_dict()
        raw=data['result']['legacy_result'];raw['slab'][0]['diameter']=4
        data['result']['slab']['sections']=deepcopy(raw['slab'])
        data['metadata']['legacy_result_sha256']=canonical_hash(raw)
        ref=self.tool.store.save(ToolResult(**data))
        with patch.object(self.tool.adapter,'generate') as generate:
            result=self.report(ref)
            self.assertFalse(result.success)
            self.assertTrue(all(e['code']=='design_check_failed' for e in result.errors))
            generate.assert_not_called()

    def test_check_timeout_bad_identity_and_false_scope_stop_before_authoring(self):
        check=self.registry.get('check_floor_design')
        with (patch.object(check.adapter,'check',side_effect=FloorCheckError('check_timeout','simulated timeout')),
              patch.object(self.tool.adapter,'generate') as generate):
            self.assertEqual(self.report().errors[0]['code'],'check_timeout');generate.assert_not_called()
        original=check._execute
        for field in ('reference','sha','coverage','summary'):
            def corrupt(data):
                payload=original(data).to_dict()
                if field=='reference':payload['result']['design_result_ref']='floor-'+'0'*32
                if field=='sha':payload['metadata']['design_sha256']='0'*64
                if field=='coverage':payload['result']['coverage']['not_checked']=[]
                if field=='summary':payload['result']['summary']['passed']-=1
                return ToolResult(**payload)
            with (patch.object(check,'_execute',side_effect=corrupt),patch.object(self.tool.adapter,'generate') as generate):
                result=self.report();self.assertFalse(result.success);generate.assert_not_called()

    def test_report_worker_timeout_start_failure_and_bad_protocol_fail_closed(self):
        # Patch only the report invocation, leaving the read-only Check worker real.
        original=subprocess.run
        def invoke_fault(fault):
            def invoke(command,**kwargs):
                if any(str(part).endswith('_report_worker.py') for part in command):
                    if isinstance(fault,Exception):raise fault
                    return SimpleNamespace(returncode=0,stdout=fault)
                return original(command,**kwargs)
            return invoke
        for fault,code in ((subprocess.TimeoutExpired('report',.01),'report_timeout'),
                           (OSError('simulated launch failure'),'report_start_failed'),
                           ('not JSON','report_protocol_error'),('{"protocol":true,"success":true}','report_protocol_error'),
                           ('{"protocol":1,"success":false}','report_generation_failed')):
            with patch('tools.floor.report_adapter.subprocess.run',side_effect=invoke_fault(fault)):
                self.assertEqual(self.report().errors[0]['code'],code)
        self.assertFalse(list((self.root/'reports').glob('*/report_manifest.json')))

    def test_completed_revision_history_is_appended_and_corruption_is_rejected(self):
        registry=ToolRegistry();register_floor_report_workflow(registry,self.root/'revision',revision=True,cad_backend=self.backend)
        request=json.loads((ROOT/'demos/revision-a.json').read_text(encoding='utf-8'))
        revised=registry.get('design_floor_with_revisions').execute(dict(project_id='qa',tool='design_floor_with_revisions',context=CONTEXT,parameters=request))
        self.assertTrue(revised.success,revised.errors)
        parameters=dict(design_result_ref=revised.result['final_reference'],revision_id=revised.result['revision_id'])
        tool=registry.get('generate_floor_report')
        envelope=dict(project_id='qa',tool=tool.name,context=CONTEXT,parameters=parameters)
        outcome=tool.execute(envelope);self.assertTrue(outcome.success,outcome.errors)
        doc=Document(next(a['path'] for a in outcome.artifacts if a['type']=='docx'))
        text='\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('附录 授权调整记录',text);self.assertIn('100 → 125 mm',text);self.assertIn('125 → 150 mm',text)
        folder=self.root/'revision/revisions'/parameters['revision_id']
        for name,mutate in [('state.json',lambda d:d.update(status='RUNNING')),
                            ('attempt-01/input.json',lambda d:d['model']['loads'].update(live_kN_m2=9)),
                            ('attempt-01/diagnosis.json',lambda d:d['result'].update(reason=None))]:
            path=folder/name;before=path.read_bytes();data=json.loads(before);mutate(data)
            path.write_text(json.dumps(data),encoding='utf-8')
            with patch.object(tool.adapter,'generate') as generate:
                self.assertEqual(tool.execute(envelope).errors[0]['code'],'report_evidence_invalid');generate.assert_not_called()
            path.write_bytes(before)

    def test_report_workflow_uses_generic_controller_and_gates_optional_cad(self):
        for fail in (False,True):
            root=self.root/('failed' if fail else 'passed');registry=ToolRegistry()
            backend=SimulatedCAD();workflow=register_floor_report_workflow(registry,root,cad=True,cad_backend=backend)
            controller=AgentController(registry,ReportRequestParser(self.parameters),AgentState(root/'agent'),[workflow])
            report=registry.get('generate_floor_report')
            with patch.object(backend,'execute',wraps=backend.execute) as cad:
                if fail:
                    with patch.object(report.adapter,'generate',side_effect=FloorReportError('report_timeout','simulated timeout')):
                        result=controller.run('explicit report request',project_id='qa')
                else:result=controller.run('explicit report request',project_id='qa')
            self.assertEqual(result['success'],not fail,result)
            self.assertEqual(cad.call_count,0 if fail else 1)
            self.assertEqual(result['steps']['cad'],'skipped' if fail else 'completed')

    def test_worker_receipt_cannot_accept_corrupt_or_missing_docx(self):
        original=subprocess.run
        for missing in (False,True):
            def fake(command,**kwargs):
                if not any(str(part).endswith('_report_worker.py') for part in command):return original(command,**kwargs)
                folder=Path(json.loads(kwargs['input'])['directory'])
                if not missing:(folder/'floor_report.docx').write_bytes(b'Not a DOCX')
                return SimpleNamespace(returncode=0,stdout=json.dumps(dict(protocol=1,success=True,
                    result=dict(docx_sha256='0'*64,paragraphs=1,tables=1,images=1))))
            with patch('tools.floor.report_adapter.subprocess.run',side_effect=fake):
                self.assertEqual(self.report().errors[0]['code'],'report_protocol_error')
        self.assertFalse(list((self.root/'reports').glob('*/report_manifest.json')))

    def test_revision_report_workflow_persists_both_steps_without_cad(self):
        registry=ToolRegistry();workflow=register_floor_report_workflow(registry,self.root/'flow',revision=True,cad_backend=self.backend)
        request=json.loads((ROOT/'demos/revision-a.json').read_text(encoding='utf-8'))
        controller=AgentController(registry,ReportRequestParser(request,True),AgentState(self.root/'flow/agent'),[workflow])
        with patch.object(self.backend,'execute',wraps=self.backend.execute) as cad:
            result=controller.run('explicit revision report request',project_id='qa')
        self.assertTrue(result['success'],result);cad.assert_not_called()
        self.assertEqual(result['steps'],dict(parse='completed',revision='completed',report='completed'))
        report=json.loads(Path(result['tool_calls'][1]['result_path']).read_text(encoding='utf-8'))
        self.assertIsNotNone(report['result']['revision_id'])
        self.assertEqual(report['result']['check_summary']['failed'],0)

    def test_manifest_write_failure_never_returns_success(self):
        from core.persistence import write_json
        def fail_manifest(path,data,**kwargs):
            if Path(path).name=='report_manifest.json':raise OSError('simulated publication failure')
            return write_json(path,data,**kwargs)
        with patch('tools.floor.report_adapter.write_json',side_effect=fail_manifest):
            result=self.report()
        self.assertFalse(result.success);self.assertEqual(result.errors[0]['code'],'report_publish_failed')
        self.assertFalse(list((self.root/'reports').glob('*/report_manifest.json')))


if __name__=='__main__':unittest.main()
