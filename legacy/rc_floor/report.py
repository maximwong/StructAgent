"""Generate editable DOCX and a PDF converted from that exact DOCX."""
from pathlib import Path
from datetime import datetime
import argparse
import json
import hashlib
import subprocess
import shutil
import sys
import csv
import os
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from engine import calculate, fmt
from figures import generate as draw_figures

ROOT=Path(__file__).resolve().parent


def set_font(run,size=10.5,bold=False):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    rpr=run._element.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
    if fonts is None:fonts=OxmlElement('w:rFonts');rpr.insert(0,fonts)
    fonts.set(qn('w:eastAsia'),'宋体')


def field(paragraph,instruction):
    run=paragraph.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');run._r.append(begin)
    ins=OxmlElement('w:instrText');ins.set(qn('xml:space'),'preserve');ins.text=' '+instruction+' ';run._r.append(ins)
    sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');run._r.append(sep)
    run=paragraph.add_run('请更新域以显示目录页码' if instruction.startswith('TOC') else '1');set_font(run,9)
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');paragraph.add_run()._r.append(end)


def para(doc,text='',style=None):
    p=doc.add_paragraph(style=style);set_font(p.add_run(text));return p


def table(doc,title,headers,rows,counter):
    cap=para(doc,f'表 {counter}  {title}');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;cap.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    width=170
    if len(headers)==2:widths=[48,122]
    elif len(headers)>=7:widths=[32]+[(width-32)/(len(headers)-1)]*(len(headers)-1)
    elif len(headers)==6:widths=[25,49,20,22,27,27]
    else:widths=[width/len(headers)]*len(headers)
    # Different six-column tables can need compact uniform numeric columns.
    if len(headers)==6 and headers[0] in ['活载布置','段']:widths=[30,28,28,28,28,28]
    # The slab moment table keeps a wide substitution column for readable numbers.
    if len(headers)==5 and headers[0]=='控制截面':widths=[26,20,24,58,42]
    for col,w in zip(t.columns,widths):col.width=Mm(w)
    for cell,w in zip(t.rows[0].cells,widths):cell.width=Mm(w)
    for cell,value in zip(t.rows[0].cells,headers):cell.text=str(value)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for cell,w,value in zip(cells,widths,row):cell.width=Mm(w);cell.text=str(value)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'5');b.set(qn('w:color'),'777777');borders.append(b)
    t._tbl.tblPr.append(borders)
    for ir,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');margins.append(e)
            cell._tc.get_or_add_tcPr().append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.0
                if ir==0:p.paragraph_format.keep_with_next=True
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if len(headers)==2 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:set_font(run,9 if len(headers)<9 else 8,ir==0)
    para(doc).paragraph_format.space_after=Pt(1)


def matrix_table(doc,title,groups,rows,counter):
    """Sample-format capacity table: sections as columns, items as rows.

    ``groups`` is ``[(label,[sub,...] or None),...]``; a group without sub-labels
    gets a single vertically merged header cell, so the same helper prints both
    the two-tier slab table (I-I / II-II bands) and the one-tier beam tables.
    """
    cap=para(doc,f'表 {counter}  {title}');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;cap.paragraph_format.keep_with_next=True
    spans=[1 if sub is None else len(sub) for _,sub in groups]
    n=sum(spans);nhead=2 if any(sub is not None for _,sub in groups) else 1
    label_w=34 if n<=5 else 30
    body=(170-label_w)/(n-1)
    widths=[label_w]+[body]*(n-1)
    t=doc.add_table(rows=nhead+len(rows),cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Mm(w)
    for row in t.rows:
        for cell,w in zip(row.cells,widths):cell.width=Mm(w)
    if nhead==2:
        col=0
        for (label,sub),span in zip(groups,spans):
            cell=t.cell(0,col)
            if span>1:cell=cell.merge(t.cell(0,col+span-1))
            cell.text=str(label)
            if sub is not None:
                for k,text in enumerate(sub):t.cell(1,col+k).text=str(text)
            else:
                # A single-level group spans both header rows, as in sample 1.
                t.cell(0,col).merge(t.cell(1,col)).text=str(label)
            col+=span
    else:
        for j,(label,_) in enumerate(groups):t.cell(0,j).text=str(label)
    for i,(label,values) in enumerate(rows):
        r=i+nhead
        t.cell(r,0).text=str(label)
        for j,value in enumerate(values):t.cell(r,j+1).text=str(value)
    for row in t.rows:
        repeat=None
        if nhead==2 and row is t.rows[0]:repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'5');b.set(qn('w:color'),'777777');borders.append(b)
    t._tbl.tblPr.append(borders)
    for ir,row in enumerate(t.rows):
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            # A merged cell is reported once per grid column, so guard against
            # writing w:tcMar twice (w:tcPr allows a single tcMar child).
            if tcpr.find(qn('w:tcMar')) is None:
                margins=OxmlElement('w:tcMar')
                for side in ['top','bottom','left','right']:
                    e=OxmlElement('w:'+side);e.set(qn('w:w'),'40');e.set(qn('w:type'),'dxa');margins.append(e)
                tcpr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1);p.paragraph_format.line_spacing=1.0
                if ir<nhead:p.paragraph_format.keep_with_next=True
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if (ir>=nhead and p is cell.paragraphs[0] and cell is row.cells[0]) else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:set_font(run,9 if ir<nhead else 8.5,ir<nhead)
    para(doc).paragraph_format.space_after=Pt(1)


def build_docx(r,figures,path):
    doc=Document();sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.left_margin=Mm(20);sec.right_margin=Mm(20);sec.top_margin=Mm(20);sec.bottom_margin=Mm(20)
    sec.header_distance=Mm(10);sec.footer_distance=Mm(10)
    for name,size in [('Normal',10.5),('Title',24),('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
        s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
        s.element.get_or_add_rPr().append(OxmlElement('w:rFonts'));s.element.rPr.rFonts.set(qn('w:eastAsia'),'黑体' if name.startswith('Heading') else '宋体')
        s.paragraph_format.line_spacing=1.15;s.paragraph_format.space_after=Pt(5)
        if name.startswith('Heading'):s.paragraph_format.space_before=Pt(12);s.paragraph_format.keep_with_next=True
    doc.styles['Normal'].paragraph_format.widow_control=True
    title=doc.add_paragraph(style='Title');title.alignment=WD_ALIGN_PARAGRAPH.CENTER;title.paragraph_format.space_before=Pt(95)
    set_font(title.add_run('单向板肋梁楼盖\n设计说明书'),24,True)
    p=r['input'];para(doc,p['project']).alignment=WD_ALIGN_PARAGRAPH.CENTER
    for label,value in [('姓名',p['report'].get('author','')),('班级',p['report'].get('class_name','')),('学号',p['report'].get('student_id','')),('日期',p['report'].get('date') or datetime.now().strftime('%Y年%m月%d日'))]:
        pr=para(doc,label+'：'+(value or '未填写'));pr.alignment=WD_ALIGN_PARAGRAPH.CENTER;pr.paragraph_format.space_after=Pt(15)
    pr=para(doc,'按样例1编排  参数化复算版');pr.alignment=WD_ALIGN_PARAGRAPH.CENTER
    para(doc,'包含板、次梁、主梁的计算过程、配筋图表及编号钢筋长度。程序内检查结果和仍需人工校核的项目见文末。')
    doc.add_page_break();doc.add_heading('目录',0);field(doc.add_paragraph(),'TOC \\o "1-2" \\h \\z \\u')
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;field(footer,'PAGE')
    tc=0;fc=0
    for chapter in r['chapters']:
        heading=doc.add_heading(chapter['title'],1);heading.paragraph_format.page_break_before=True
        for item in chapter['items']:
            kind=item['type']
            if kind=='text':para(doc,item['text'])
            elif kind=='heading':doc.add_heading(item['text'],2)
            elif kind=='equation':
                pp=para(doc,item['label']+'：'+item['formula']);pp.paragraph_format.keep_with_next=True;pp.paragraph_format.space_after=Pt(1)
                pp=para(doc,'代入 '+item['substitution']+' = '+fmt(item['result'],5)+' '+item['unit']);pp.paragraph_format.left_indent=Mm(6);pp.paragraph_format.space_after=Pt(6)
            elif kind=='table':tc+=1;table(doc,item['title'],item['headers'],item['rows'],tc)
            elif kind=='matrix':tc+=1;matrix_table(doc,item['title'],item['groups'],item['rows'],tc)
            elif kind=='figure':
                files=figures[item['key']]
                for i,f in enumerate(files):
                    fc+=1
                    from PIL import Image
                    with Image.open(f) as im:w,h=im.size
                    # A figure is at most 180mm tall so it can share a page.
                    width=170 if item['key'].endswith('_shapes') else min(170,180*w/h)
                    if item['key'].endswith('_shapes') and h/w*width>180:raise ValueError('固定比例钢筋图超过页面高度')
                    pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.keep_with_next=True
                    pp.add_run().add_picture(str(f),width=Mm(width))
                    caption=item['caption']
                    if item['key']=='slab_bars':
                        caption=['五跨代表板带配筋（非楼盖全部跨数）','板墙端与内支座配筋节点','板分布筋与墙角双向构造筋平面局部'][i]
                    cap=para(doc,f'图 {fc}  '+caption+(f'（{i+1}/{len(files)}）' if len(files)>1 else ''));cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.core_properties.title=p['project'];doc.core_properties.subject='样例1同结构参数化计算';doc.core_properties.author=p['report'].get('author','')
    # Some bundled Word templates carry an accent border in the Title style.
    for tree in [doc.styles.element,doc._element]:
        for node in list(tree.iter(qn('w:pBdr'))):node.getparent().remove(node)
    doc.save(path)


def convert(docx,pdf,timeout=180):
    if os.name=='nt':
        cmd=['powershell.exe','-NoProfile','-ExecutionPolicy','Bypass','-File',str(ROOT/'word_to_pdf.ps1'),'-InputDocx',str(docx.resolve()),'-OutputPdf',str(pdf.resolve())]
        try:
            result=subprocess.run(cmd,capture_output=True,timeout=timeout,creationflags=subprocess.CREATE_NO_WINDOW)
            if result.returncode==0 and pdf.exists():return 'Microsoft Word'
            word_error=result.stderr.decode('utf-8',errors='replace')[-1500:]
        except (OSError,subprocess.TimeoutExpired) as e:word_error=str(e)
    else:word_error='非Windows系统'
    soffice=shutil.which('soffice') or shutil.which('soffice.exe')
    if not soffice and os.name=='nt':
        candidate=Path(os.environ.get('PROGRAMFILES','C:/Program Files'))/'LibreOffice/program/soffice.exe'
        if candidate.exists():soffice=str(candidate)
    if soffice:
        import tempfile
        with tempfile.TemporaryDirectory(prefix='rcfloor_lo_') as tmp:
            cmd=[soffice,'-env:UserInstallation='+Path(tmp).as_uri(),'--headless','--convert-to','pdf','--outdir',str(pdf.parent.resolve()),str(docx.resolve())]
            subprocess.run(cmd,check=True,capture_output=True,timeout=timeout)
        if pdf.exists():return 'LibreOffice'
    raise RuntimeError('Word转PDF未完成。请安装并打开一次桌面Word，或安装LibreOffice；DOCX已保留。\n'+word_error)


def generate(config,out,docx_only=False):
    from workflow import generate as run
    return run(config,out,docx_only)


def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('input',nargs='?',default=str(ROOT/'sample1.json'));ap.add_argument('--out');ap.add_argument('--docx-only',action='store_true');args=ap.parse_args()
    out=Path(args.out) if args.out else ROOT/'generated'/datetime.now().strftime('%Y%m%d_%H%M%S_%f')
    try:
        result=generate(args.input,out,args.docx_only);print('输出目录：'+str(result.resolve()))
        state=json.loads((result/'STATUS.json').read_text(encoding='utf-8'))
        if state['state'] not in ('SUCCESS','DOCX_READY'):print('部分输出未完成，详见问题清单。',file=sys.stderr);return 2
    except Exception as e:print('生成失败：'+str(e),file=sys.stderr);return 1
    return 0


if __name__=='__main__':raise SystemExit(main())
