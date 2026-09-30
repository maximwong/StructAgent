import unittest
import sys
import json
import copy
import tempfile
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from engine import calculate, capacity, fmt, spec
from legacy_core import flexure, run as core_run
from beam_solver import beam, moment
from app import parse_value, LABELS, TABS, HIDDEN


class ReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.p=json.loads((ROOT/'sample1.json').read_text(encoding='utf-8'));cls.r=calculate(cls.p)

    def test_single_span_analytical(self):
        r=beam([6],[[10,10]])
        self.assertAlmostEqual(r['reactions'][0],10);self.assertAlmostEqual(moment(r,2),20)

    def test_sample_loads_and_coefficients(self):
        r=self.r
        self.assertAlmostEqual(r['loads']['slab_gk'],2.592)
        self.assertAlmostEqual(r['loads']['secondary_gk'],6.90688)
        self.assertAlmostEqual(r['slab'][0]['M_kNm'],(2.592*1.3+6.5*1.5)*1.82**2/11)
        self.assertAlmostEqual(r['slab'][1]['M_kNm'],-r['slab'][0]['M_kNm'])
        self.assertAlmostEqual(r['secondary'][3]['M_kNm'],-(6.90688*1.3+19.5)*5.75**2/14)

    def test_each_case_equilibrium(self):
        for c in self.r['envelope']['cases']:
            self.assertLess(abs(c['equilibrium_error']),1e-8)
            for a,b in zip(c['segments'],c['segments'][1:]):self.assertAlmostEqual(a['m1'],b['m0'],places=7)

    def test_mirror_cases(self):
        cases={x['pattern']:x for x in self.r['envelope']['cases']};L=sum(self.r['spans']['main_m'])
        for pattern,c in cases.items():
            for x in [0,1,5,8,12,17]:self.assertAlmostEqual(moment(c,x),moment(cases[pattern[::-1]],L-x),places=7)

    def test_capacity_does_not_use_area_ratio(self):
        c=flexure(80,250,500,460,14.3,360,1.43);c.update(As_provided=c['As_required']*1.2,M_kNm=80)
        p=copy.deepcopy(self.p);p['main'].update(b_mm=250,h_mm=500)
        _,mu=capacity(c,p,'main');self.assertGreater(mu,80);self.assertLess(mu,96)

    def test_face_moment_same_case(self):
        p=copy.deepcopy(self.p);p['main']['support_moment']='axis';r=core_run(p)
        self.assertGreater(abs(r['main'][1]['M_kNm']),abs(self.r['main'][1]['M_kNm']))

    def test_lap_congestion_rejected(self):
        p=copy.deepcopy(self.p);p['main']['support_moment']='axis'
        with self.assertRaisesRegex(ValueError,'搭接'):calculate(p)

    def test_material_regions_and_coverage(self):
        for c in self.r['material_checks']:self.assertGreaterEqual(c['capacity_kNm']+1e-5,c['demand_kNm'])
        self.assertTrue(self.r['material_checks'])

    def test_additional_stirrups_even_and_sufficient(self):
        s=self.r['suspension'];self.assertEqual(s['count']%2,0);self.assertEqual(s['count'],s['per_side']*2)
        self.assertGreaterEqual(s['capacity_kN'],s['F_kN']);self.assertLessEqual(s['spacing_mm']*s['count'],s['range_mm'])

    def test_lengths_sum_and_stock(self):
        for b in self.r['bars']:
            self.assertAlmostEqual(b['length_mm'],sum(v for _,v in b['segments']))
            self.assertGreaterEqual(b['rounded_mm'],b['length_mm']);self.assertLessEqual(b['rounded_mm'],self.p['report']['max_stock_mm'])

    def test_input_unchanged(self):
        old=copy.deepcopy(self.p);calculate(self.p);self.assertEqual(old,self.p)

    def test_load_and_height_changes_propagate(self):
        r=calculate(json.loads((ROOT/'changed.json').read_text(encoding='utf-8')))
        self.assertNotEqual(r['main'][0]['M_kNm'],self.r['main'][0]['M_kNm'])
        self.assertNotEqual(r['main'][0]['As_provided'],self.r['main'][0]['As_provided'])
        self.assertNotEqual(r['bars'][-1]['length_mm'],self.r['bars'][-1]['length_mm'])
        self.assertNotEqual(r['chapters'],self.r['chapters'])

    def test_outside_scope_rejected(self):
        for key,value in [('secondary_axis_spans_mm',[6000]*4),('main_axis_spans_mm',[6000]*4),('secondary_spacing_mm',1900)]:
            p=copy.deepcopy(self.p);p['geometry'][key]=value
            with self.assertRaises(ValueError):calculate(p)

    def test_high_load_and_nan_rejected(self):
        for v in [30,float('nan'),-1]:
            p=copy.deepcopy(self.p);p['loads']['live_kN_m2']=v
            with self.assertRaises(ValueError):calculate(p)

    def test_short_stock_rejected(self):
        p=copy.deepcopy(self.p);p['report']['max_stock_mm']=1000
        with self.assertRaises(ValueError):calculate(p)

    def test_parameter_parser(self):
        self.assertEqual(parse_value('6000，6000,6000',[6000],'跨度'),[6000,6000,6000])
        with self.assertRaises(ValueError):parse_value('2.5',2,'排数')

    def test_all_visible_parameters_translated(self):
        for _,groups in TABS:
            for group in groups:
                v=self.p[group]
                for key in (v if isinstance(v,dict) else [group]):self.assertTrue(key in LABELS or key in HIDDEN,key)

    def test_all_chapter_figures_exist(self):
        from figures import generate
        with tempfile.TemporaryDirectory() as tmp:
            figs=generate(self.r,tmp)
            for chapter in self.r['chapters']:
                for i in chapter['items']:
                    if i['type']=='figure':
                        for p in figs[i['key']]:self.assertTrue(p.exists())

    def test_slab_coefficients_stored_once_and_consistent(self):
        r=self.r;g=r['loads']['slab_g'];q=r['loads']['slab_q']
        self.assertEqual([(x['alpha_num'],x['alpha_den']) for x in r['slab'][:4]],[(1,11),(-1,11),(1,16),(-1,14)])
        for x,le in zip(r['slab'][:4],[r['spans']['slab_edge_mm']]*2+[r['spans']['slab_inner_mm']]*2):
            self.assertAlmostEqual(x['span_mm'],le)
            self.assertAlmostEqual(x['alpha_num']/x['alpha_den']*(g+q)*(le/1000)**2,x['M_kNm'],places=9)

    def test_coefficient_diagram_reads_results_and_matches_table(self):
        from figures import slab_coeff_layout
        r=self.r;Ls,mids,sups=slab_coeff_layout(r)
        self.assertEqual(Ls,[r['spans']['slab_edge_mm']]+[r['spans']['slab_inner_mm']]*3+[r['spans']['slab_edge_mm']])
        self.assertEqual([round(x['alpha_num']/x['alpha_den'],9) for x in mids],[round(1/11,9)]+[round(1/16,9)]*3+[round(1/11,9)])
        self.assertEqual([round(x['alpha_num']/x['alpha_den'],9) for x in sups],[round(-1/11,9),round(-1/14,9),round(-1/14,9),round(-1/11,9)])
        # Diagram ordinates are the very rows the report table prints.
        for x in mids+sups:self.assertIn(x['M_kNm'],[y['M_kNm'] for y in r['slab'][:4]])
        items=[i for ch in r['chapters'] for i in ch['items']]
        keys=[i.get('key') or i.get('title') or i.get('label','') for i in items]
        fig=keys.index('slab_moment_coefficients');tbl=keys.index('板弯矩计算表 1m宽板带')
        method=next(i for i,k in enumerate(keys) if k=='跨度差')
        self.assertLess(method,fig);self.assertLess(fig,tbl)
        rows=items[tbl]['rows']
        self.assertEqual([row[0] for row in rows],['边跨中','第一内支座','中间跨中','中间支座'])
        self.assertEqual([row[1] for row in rows],['1/11','-1/11','1/16','-1/14'])
        for row,x in zip(rows,r['slab'][:4]):
            self.assertEqual(row[4],('+' if x['M_kNm']>=0 else '')+fmt(x['M_kNm'],5))

    def test_coefficients_stable_and_values_update_with_inputs(self):
        r=calculate(json.loads((ROOT/'changed.json').read_text(encoding='utf-8')))
        self.assertEqual([(x['alpha_num'],x['alpha_den']) for x in r['slab'][:4]],[(1,11),(-1,11),(1,16),(-1,14)])
        self.assertNotEqual(r['slab'][0]['M_kNm'],self.r['slab'][0]['M_kNm'])
        self.assertNotEqual(r['loads']['slab_q'],self.r['loads']['slab_q'])

    def test_out_of_scope_never_draws_coefficient_diagram(self):
        import copy
        from figures import slab_coeff_layout
        r=copy.deepcopy(self.r);r['loads']['slab_q']=r['loads']['slab_g']*3.5
        with self.assertRaisesRegex(ValueError,'系数法'):slab_coeff_layout(r)

    def test_failed_generation_status(self):
        from report import generate
        p=copy.deepcopy(self.p);p['loads']['live_kN_m2']=100
        with tempfile.TemporaryDirectory() as tmp:
            source=Path(tmp)/'bad.json';source.write_text(json.dumps(p),encoding='utf-8')
            with self.assertRaises(ValueError):generate(source,Path(tmp)/'out',True)
            self.assertEqual(json.loads((Path(tmp)/'out/STATUS.json').read_text(encoding='utf-8'))['state'],'FAILED')
            self.assertFalse((Path(tmp)/'out/设计说明书.docx').exists())

    # --- Detailing text and sample-format capacity tables (改进汇总 requirements) ---

    def _items(self):
        return [i for ch in self.r['chapters'] for i in ch['items']]

    def _table(self,title):
        return next(i for i in self._items() if i.get('title')==title)

    def test_body_tables_keep_sample_numbering(self):
        tables=[i['title'] for i in self._items() if i['type'] in ('table','matrix')]
        self.assertEqual(tables[:10],['板弯矩计算表 1m宽板带','板截面承载力计算','次梁弯矩计算表','次梁剪力计算表',
            '次梁正截面承载力计算','次梁斜截面承载力计算','主梁弯矩计算表 kN·m','主梁剪力计算表 kN',
            '主梁正截面承载力计算','主梁斜截面承载力计算'])
        # Length summaries and the reaction table moved to the appendix, so the
        # body still reproduces the sample-1 table numbering.
        appendix=next(ch for ch in self.r['chapters'] if ch['title'].startswith('附录'))
        titles=[i['title'] for i in appendix['items'] if i['type'] in ('table','matrix')]
        self.assertIn('板编号钢筋长度汇总',titles);self.assertIn('主梁各工况反力与平衡',titles)

    def test_secondary_coefficients_defined_once(self):
        for x,(num,den),le in zip(self.r['secondary'],[(1,11),(-1,11),(1,16),(-1,14)],
                                  [self.r['spans']['secondary_edge_mm']]*2+[self.r['spans']['secondary_inner_mm']]*2):
            self.assertEqual((x['alpha_num'],x['alpha_den']),(num,den));self.assertAlmostEqual(x['span_mm'],le)
        g=self.r['loads']['secondary_g'];q=self.r['loads']['secondary_q']
        for x in self.r['secondary']:self.assertAlmostEqual(x['alpha']*(g+q)*(x['span_mm']/1000)**2,x['M_kNm'],places=9)
        for x,(num,den) in zip(self.r['secondary_shear'],[(45,100),(6,10),(55,100),(55,100)]):
            self.assertEqual((x['beta_num'],x['beta_den']),(num,den))
            self.assertAlmostEqual(x['beta']*(g+q)*x['span_mm']/1000,x['V_kN'],places=9)

    def test_secondary_force_tables_read_stored_coefficients(self):
        moment=self._table('次梁弯矩计算表');shear=self._table('次梁剪力计算表')
        self.assertEqual([g[0] for g in moment['groups']],['截面','边跨中','第一内支座','中间跨中','中间支座'])
        self.assertEqual([g[0] for g in shear['groups']][1:],[x['name'] for x in self.r['secondary_shear']])
        self.assertEqual(moment['rows'][0][1],['1/11','-1/11','1/16','-1/14'])
        for cell,x in zip(moment['rows'][1][1],self.r['secondary']):self.assertIn(fmt(x['M_kNm'],3),cell)
        self.assertEqual(shear['rows'][0][1],['0.45','0.6','0.55','0.55'])
        for cell,x in zip(shear['rows'][1][1],self.r['secondary_shear']):self.assertIn(fmt(x['V_kN'],3),cell)

    def test_slab_capacity_table_keeps_two_bands(self):
        t=self._table('板截面承载力计算')
        self.assertEqual([g[0] for g in t['groups']],['截面','边跨中','第一内支座','中间跨中','中间支座'])
        self.assertEqual(t['groups'][3][1],['Ⅰ-Ⅰ板带','Ⅱ-Ⅱ板带']);self.assertEqual(t['groups'][4][1],['Ⅰ-Ⅰ板带','Ⅱ-Ⅱ板带'])
        for _,values in t['rows']:self.assertEqual(len(values),6)
        rows=self.r['slab'];order=[rows[0],rows[1],rows[2],rows[4],rows[3],rows[5]]
        pick=next(v for l,v in t['rows'] if l.startswith('实配面积'))
        self.assertEqual(pick,[fmt(x['As_provided'],1) for x in order])
        # The middle band really is the arch-reduced row, not a copy of the edge band.
        self.assertLess(abs(rows[4]['M_kNm']),abs(rows[2]['M_kNm']))

    def test_beam_capacity_tables_cover_every_section(self):
        for title,rows in [('次梁正截面承载力计算',self.r['secondary']),
                           ('次梁斜截面承载力计算',self.r['secondary_shear']),
                           ('主梁正截面承载力计算',self.r['main']),
                           ('主梁斜截面承载力计算',self.r['main_shear'])]:
            t=self._table(title)
            self.assertEqual([g[0] for g in t['groups']][1:],[x['name'] for x in rows])
            for _,values in t['rows']:self.assertEqual(len(values),len(rows))
        main=self._table('主梁正截面承载力计算');labels=[l for l,_ in main['rows']]
        self.assertIn('其中轴线弯矩 M轴 (kN·m)',labels);self.assertIn('V0·b柱/2 (kN·m)',labels)
        sup=self.r['support']
        self.assertAlmostEqual(sup['reduction_kNm'],sup['v0_kN']*sup['b_mm']/2000,places=9)
        self.assertEqual(main['rows'][labels.index('V0·b柱/2 (kN·m)')][1][1],fmt(sup['reduction_kNm'],3))
        pick=next(v for l,v in main['rows'] if l.startswith('选配钢筋'))
        self.assertEqual(pick,[spec(x) for x in self.r['main']])

    def test_detailing_narrative_present(self):
        text='\n'.join(i['text'] for i in self._items() if i['type']=='text')
        for token in ['Ⅰ-Ⅰ板带是指板的边带','Ⅱ-Ⅱ板带在Ⅰ-Ⅰ板带的基础上下降20%','本设计采用分离式配筋方案',
                      '受力钢筋的间距不超过200mm','分布钢筋的截面面积不宜小于单位宽度上受力钢筋截面面积的15%',
                      '次梁在跨中按T形截面计算','主梁的中间支座按铰支座考虑','要求梁柱线刚度之比超过4',
                      '次梁的边跨等于净跨加次梁在边支座上的支承长度的一半','下部受拉纵筋不宜在跨中截面截断']:
            self.assertIn(token,text)

    def test_matrix_table_renders_merged_two_tier_header(self):
        from report import build_docx
        from figures import generate
        from docx import Document as Docx
        with tempfile.TemporaryDirectory() as tmp:
            figs=generate(self.r,tmp);path=Path(tmp)/'r.docx';build_docx(self.r,figs,path)
            d=Docx(str(path))
            t=next(tb for tb in d.tables if len(tb.columns)==7 and tb.rows[0].cells[0].text.strip()=='截面')
            self.assertEqual(t.rows[0].cells[3].text.strip(),'中间跨中')
            self.assertEqual(t.rows[1].cells[3].text.strip(),'Ⅰ-Ⅰ板带');self.assertEqual(t.rows[1].cells[4].text.strip(),'Ⅱ-Ⅱ板带')
            self.assertIs(t.cell(0,0)._tc,t.cell(1,0)._tc)
            self.assertIs(t.cell(0,1)._tc,t.cell(1,1)._tc)


if __name__=='__main__':unittest.main()
