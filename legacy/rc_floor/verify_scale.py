"""Internal 6000/4000 mm physical-scale fixture, including real Word PDF export."""
from pathlib import Path
import json,copy
from docx import Document
from docx.shared import Mm
from detail_geometry import enrich_details
from scaled_figures import shapes


def prepare(root):
    out=Path(root)/'verification';out.mkdir(exist_ok=True)
    r=json.loads((Path(root)/'examples/sample1/results.json').read_text(encoding='utf-8'))
    r['bars']=[b for b in r['bars'] if b['member']!='slab']
    for i,L in enumerate([6000,4000,600,6000,4000]):
        r['bars'].append(dict(mark=f'T{i+1}',member='slab',use='比例验收直筋',diameter=12,count=1,pieces=1,segments=[('直长',L)],length_mm=L,rounded_mm=L,shape='straight',notes=''))
    enrich_details(r);doc=Document();sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297);sec.left_margin=sec.right_margin=Mm(20)
    for pg in range(2):
        if pg:doc.add_page_break()
        f=out/f'ratio_{pg}.png';shapes(r,'slab',pg).save(f);doc.add_paragraph().add_run().add_picture(str(f),width=Mm(170))
    doc.save(out/'6000_4000.docx');return out


def verify_pdf(path):
    """Measure the black primary lines in native Word-exported pages, not source images."""
    import pypdfium2 as pdf
    document=pdf.PdfDocument(str(path));measured=[];page_results=[]
    for page in document:
        im=page.render(scale=3).to_pil().convert('RGB');runs=[]
        # Horizontal primary lines are pure black and much longer than any letters.
        for y in range(im.height):
            best=0;start=None
            for x in range(im.width):
                dark=max(im.getpixel((x,y)))<50
                if dark and start is None:start=x
                if (not dark or x==im.width-1) and start is not None:
                    best=max(best,x-start);start=None
            if best>im.width*.25 and (not runs or y-runs[-1][0]>5):runs.append((y,best))
        # Grey comparison lines deliberately fail the pure-black threshold.
        lengths=[px/3*25.4/72 for _,px in runs]
        page_results.append(lengths)
        if len(lengths)<2:raise AssertionError('PDF中未识别两条验收直筋')
        a,b=lengths[:2]
        if abs(a-120)>.7 or abs(b-80)>.7 or abs(a/b-1.5)>.012:raise AssertionError('PDF实际显示比例不符：'+str(lengths))
        measured.append(a/b)
    if len(document)!=2:raise AssertionError('跨页验收应为两页')
    return dict(state='SUCCESS',physical_lengths_mm=page_results,ratios=measured,expected_mm=[120,80],scale='1:50',method='native Word PDF raster at 3x; pixel tolerance 0.7mm')


def run(root):
    from report import convert
    out=prepare(root);convert(out/'6000_4000.docx',out/'6000_4000.pdf');result=verify_pdf(out/'6000_4000.pdf')
    (out/'scale_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');return result
