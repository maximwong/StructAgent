"""Append traceable check evidence to the existing calculation-book layout."""

from docx import Document
from docx.shared import Mm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE


def augment(path, evidence):
    # This function runs only in the isolated report worker.
    from report import para, set_font, fmt
    doc = Document(path)
    # Define Word's TOC styles before field updates to avoid a near-empty
    # overflow page when the new audit headings are added to the old book.
    for level in (1,2):
        name='toc '+str(level)
        style=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH,builtin=True)
        style.style_id='TOC'+str(level)
        style.base_style=doc.styles['Normal']
        style.font.size=Pt(9.5)
        style.paragraph_format.line_spacing=1
        style.paragraph_format.space_before=Pt(0)
        style.paragraph_format.space_after=Pt(2)
        style.paragraph_format.left_indent=Mm(4 if level==2 else 0)
    for text in (
        '设计引用：' + evidence['design_result_ref'],
        '设计结果 SHA256：' + evidence['design_sha256'],
        '本计算书对应已保存方案，含原程序计算过程及独立截面校核附录。校核通过范围见附录，不能据此认定工程全项审查通过。',
    ):
        p = doc.add_paragraph()
        set_font(p.add_run(text), 10.5)
        # Put provenance on the existing cover, before its page break.
        first_break = next(p for p in doc.paragraphs if p._p.xpath('.//w:br[@w:type="page"]'))
        first_break._p.addprevious(p._p)
    heading = doc.add_heading('附录 独立截面校核记录', 1)
    heading.paragraph_format.page_break_before = True
    report = evidence['check']['result']
    summary = report['summary']
    para(doc, f"校核状态 {report['status']}；通过 {summary['passed']} 项，共 {summary['total']} 项；未通过 {summary['failed']} 项。")
    coverage = report['coverage']
    para(doc, '已复核：' + '；'.join(coverage['recomputed']))
    para(doc, '需求来源：' + coverage['demand_source'])
    para(doc, '未覆盖：' + '；'.join(coverage['not_checked']))

    def audit_table(headers, rows, widths):
        table = doc.add_table(rows=1, cols=len(headers))
        table.autofit = False
        for col, width in zip(table.columns, widths): col.width = Mm(width)
        for cell, value in zip(table.rows[0].cells, headers): cell.text = value
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for row in rows:
            for cell, value in zip(table.add_row().cells, row): cell.text = str(value)
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell, width in zip(row.cells, widths):
                cell.width = Mm(width)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = False
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.space_before = Pt(2)
                    for run in paragraph.runs: set_font(run, 9)
        para(doc)

    names = {'slab': '板', 'secondary': '次梁', 'main': '主梁'}
    for member, title in names.items():
        doc.add_heading(title + '逐项校核', 2)
        checks = [item for item in report['checks'] if item['member'] == member]
        audit_table(['截面及检查', '实际值', '比较', '限值', '单位', '结果'],
            [[str(item['section'] + 1) + ' ' + item['label'], fmt(item['actual'], 5), item['relation'],
              fmt(item['limit'], 5), item['unit'], '通过' if item['passed'] else '未通过'] for item in checks],
            [62, 28, 14, 28, 20, 18])
        para(doc, '复核函数：' + '；'.join(dict.fromkeys(item['source'] for item in checks)))
    history = evidence.get('history')
    if history is not None:
        heading = doc.add_heading('附录 授权调整记录', 1)
        heading.paragraph_format.page_break_before = True
        para(doc, '调整记录：' + history['revision_id'])
        para(doc, '每轮仅修改用户授权的下一尺寸候选，重新完整设计和校核；失败方案不作为最终出图方案。')
        for record in history['records']:
            doc.add_heading('第 ' + str(record['index']) + ' 轮', 2)
            para(doc, '设计状态：' + record['design_status'])
            for error in record['errors']:
                label={'design_rejected':'设计未通过','design_check_failed':'校核未通过'}.get(error['code'],error['code'])
                para(doc, label + '：' + error['message'].replace('secondary/','次梁/').replace('main/','主梁/'))
            if record['change'] is not None:
                change = record['change']
                label={'model.slab.h_mm':'板厚','model.secondary.b_mm':'次梁宽度','model.secondary.h_mm':'次梁高度',
                       'model.main.b_mm':'主梁宽度','model.main.h_mm':'主梁高度'}[change['field']]
                para(doc, label + '：' + str(change['before']) + ' → ' + str(change['after']) + ' mm')
            para(doc, '依据：' + record['reason'])
    # Apply consistent, light borders without changing old calculation content.
    for table in doc.tables:
        borders = table._tbl.tblPr.find(qn('w:tblBorders'))
        if borders is None:
            borders = OxmlElement('w:tblBorders'); table._tbl.tblPr.append(borders)
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            node = borders.find(qn('w:' + side))
            if node is None: node = OxmlElement('w:' + side); borders.append(node)
            for attr, value in [('val', 'single'), ('sz', '5'), ('color', 'D9D9D9')]: node.set(qn('w:' + attr), value)
        for cell in table.rows[0].cells:
            shading = cell._tc.get_or_add_tcPr().find(qn('w:shd'))
            if shading is None: shading = OxmlElement('w:shd'); cell._tc.get_or_add_tcPr().append(shading)
            shading.set(qn('w:fill'), 'EAF1F8')
    doc.core_properties.last_modified_by = 'StructAgent'
    doc.save(path)
