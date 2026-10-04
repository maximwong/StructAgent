"""Bounded, isolated DOCX generation from immutable design and check evidence."""

import hashlib
import json
import math
import os
from pathlib import Path
import subprocess
import sys
from uuid import uuid4
from zipfile import ZipFile

from docx import Document

from core.persistence import write_json
from core.validation import ensure_json


class FloorReportError(ValueError):
    def __init__(self, code, message):
        self.code = code
        super().__init__(message)


class FloorReportAdapter:
    def __init__(self, root, *, timeout_seconds=60):
        if type(timeout_seconds) not in (int, float) or not math.isfinite(timeout_seconds) or timeout_seconds <= 0:
            raise ValueError('Report timeout must be positive and finite.')
        self.root, self.timeout = Path(root).resolve(), timeout_seconds

    def generate(self, design, evidence):
        report_id = 'report-' + uuid4().hex
        directory = self.root/report_id
        directory.mkdir(parents=True, exist_ok=False)
        write_json(directory/'input.json', evidence, exclusive=True)
        worker = Path(__file__).with_name('_report_worker.py')
        request = dict(directory=str(directory), legacy_result=design.result['legacy_result'], evidence=evidence)
        try:
            process = subprocess.run([sys.executable, '-I', str(worker)],
                input=json.dumps(request, ensure_ascii=False, allow_nan=False), capture_output=True, text=True,
                encoding='utf-8', cwd=str(worker.parents[2]), timeout=self.timeout,
                creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
        except subprocess.TimeoutExpired:
            raise FloorReportError('report_timeout', '计算书生成超时；保留本次记录，不自动重试。') from None
        except (OSError, UnicodeError):
            raise FloorReportError('report_start_failed', '无法启动或读取计算书生成进程。') from None
        try:
            payload = json.loads(process.stdout); ensure_json(payload)
            if (process.returncode != 0 or type(payload['protocol']) is not int or payload['protocol'] != 1
                    or type(payload['success']) is not bool): raise ValueError('Invalid report protocol.')
            if not payload['success']:
                raise FloorReportError('report_generation_failed', '计算书生成失败，不能确认本次报告可用。')
            path = directory/'floor_report.docx'
            with ZipFile(path) as archive:
                if archive.testzip() is not None or not {'[Content_Types].xml', 'word/document.xml'} <= set(archive.namelist()):
                    raise ValueError('Invalid DOCX package.')
            checksum = hashlib.sha256(path.read_bytes()).hexdigest()
            result = payload['result']
            if (result['docx_sha256'] != checksum or any(type(result[key]) is not int or result[key] <= 0
                    for key in ('paragraphs', 'tables', 'images'))): raise ValueError('Invalid report receipt.')
            document = Document(path)
            if (result['paragraphs'],result['tables'],result['images']) != (len(document.paragraphs),len(document.tables),len(document.inline_shapes)):
                raise ValueError('Report receipt disagrees with document contents.')
        except FloorReportError: raise
        except Exception:
            raise FloorReportError('report_protocol_error', '计算书文件或生成回执不完整，已停止后续执行。') from None
        manifest = dict(report_id=report_id, status='completed', format='docx', **result,
                        evidence_sha256=hashlib.sha256((directory/'input.json').read_bytes()).hexdigest(),
                        design_result_ref=evidence['design_result_ref'], design_sha256=evidence['design_sha256'],
                        check_summary=evidence['check']['result']['summary'],
                        revision_id=evidence['history']['revision_id'] if evidence['history'] else None)
        write_json(directory/'report_manifest.json', manifest, exclusive=True)
        return dict(result=manifest, artifacts=[dict(type='docx', path=str(path)),
            dict(type='report_evidence', path=str(directory/'input.json')),
            dict(type='report_manifest', path=str(directory/'report_manifest.json'))],
            metadata=dict(project_id=design.metadata['project_id'], recovery_required=False))
