"""Render saved calculation content only; no design or CAD execution."""

from contextlib import redirect_stdout
import hashlib
import json
from pathlib import Path
import sys


def main():
    for stream in (sys.stdin, sys.stdout, sys.stderr): stream.reconfigure(encoding='utf-8')
    root = Path(__file__).resolve().parents[2]
    sys.path.insert(0, str(root)); sys.path.insert(0, str(root/'legacy/rc_floor'))
    try:
        request = json.load(sys.stdin)
        with redirect_stdout(sys.stderr):
            import engine, legacy_core, report
            from figures import generate
            from tools.floor.report_builder import augment
            def forbidden(*args, **kwargs): raise RuntimeError('Report cannot redesign a saved scheme.')
            engine.calculate = legacy_core.run = report.calculate = forbidden
            directory = Path(request['directory'])
            path = directory/'floor_report.docx'
            raw = request['legacy_result']
            figures = generate(raw, directory/'figures')
            report.build_docx(raw, figures, path)
            augment(path, request['evidence'])
            from docx import Document
            document = Document(path)
            result = dict(docx_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                          paragraphs=len(document.paragraphs), tables=len(document.tables), images=len(document.inline_shapes))
        payload = dict(protocol=1, success=True, result=result)
    except Exception:
        payload = dict(protocol=1, success=False)
    sys.stdout.write(json.dumps(payload, ensure_ascii=False, allow_nan=False))


if __name__ == '__main__': main()
