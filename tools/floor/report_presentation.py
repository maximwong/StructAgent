"""Read-only floor report provenance and file checks for the local UI."""

import hashlib
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from core import ToolResult
from core.validation import make_validator, validate_json
from .check_adapter import COVERAGE, validate_report
from .check_tool import OUTPUT as CHECK_OUTPUT
from .design_store import FloorDesignStore
from .report_tool import OUTPUT


def floor_report_source(root, snapshot, read_result):
    if (snapshot.get('status') in ('queued', 'running', 'interrupted', 'recovery_required')
            or snapshot.get('external_started')
            or snapshot.get('steps', {}).get('design') != 'completed'
            or snapshot.get('steps', {}).get('check') != 'completed'):
        raise ValueError('A terminal design with a passing check is required.')
    def one(name):
        calls = [c for c in snapshot.get('tool_calls', []) if c.get('tool') == name]
        if len(calls) != 1 or calls[0].get('status') != 'completed':
            raise ValueError('Missing or duplicated source tool call.')
        return read_result(calls[0], snapshot)
    design = one('design_floor_system')
    checked = one('check_floor_design')
    reference = design.metadata['design_result_ref']
    stored = FloorDesignStore(Path(root)/'designs').load(reference, project_id=snapshot['project_id'])
    published = design.to_dict()
    published['metadata'].pop('design_result_ref',None)  # Added after immutable store publication.
    if stored.to_dict() != published:
        raise ValueError('Published design and stored reference differ.')
    _check(checked, snapshot['project_id'], reference, stored.metadata['legacy_result_sha256'])
    return dict(request=dict(project_id=snapshot['project_id'], tool='generate_floor_report',
        context=dict(unit_system='SI', design_code='GB'), parameters=dict(design_result_ref=reference)),
        design_sha256=stored.metadata['legacy_result_sha256'])


def _check(payload, project, reference, checksum):
    validate_json(payload.result, make_validator(CHECK_OUTPUT))
    validate_report({key:payload.result[key] for key in ('status','checks','summary')})
    if (not payload.success or payload.errors or payload.tool != 'check_floor_design' or payload.version != '1.0.0'
            or payload.result['status'] != 'PASS' or payload.result['coverage'] != COVERAGE
            or payload.result['design_result_ref'] != reference
            or payload.metadata.get('project_id') != project
            or payload.metadata.get('design_result_ref') != reference
            or payload.metadata.get('design_sha256') != checksum
            or payload.metadata.get('scope_version') != 'section-check-v1'):
        raise ValueError('Check provenance differs from source design.')


def validate_floor_report(root, payload, source):
    result = payload.result
    request = source['request']
    validate_json(result, make_validator(OUTPUT))
    if (not payload.success or payload.tool != request['tool'] or payload.version != '1.0.0'
            or payload.errors or payload.metadata.get('project_id') != request['project_id']
            or payload.metadata.get('recovery_required') is not False
            or result['design_result_ref'] != request['parameters']['design_result_ref']
            or result['design_sha256'] != source['design_sha256'] or result['revision_id'] is not None):
        raise ValueError('Report receipt provenance differs.')
    directory = Path(root).resolve()/'reports'/result['report_id']
    expected = {'docx':directory/'floor_report.docx', 'report_evidence':directory/'input.json',
                'report_manifest':directory/'report_manifest.json'}
    if len(payload.artifacts) != 3 or {a['type'] for a in payload.artifacts} != set(expected):
        raise ValueError('Report artifacts differ.')
    for item in payload.artifacts:
        path = expected[item['type']]
        if (Path(item['path']).resolve() != path or path.resolve() != path
                or not path.is_file() or not 0 < path.stat().st_size <= 64*1024*1024):
            raise ValueError('Report path is redirected, missing or oversized.')
    body = expected['docx'].read_bytes()
    evidence_bytes = expected['report_evidence'].read_bytes()
    evidence = json.loads(evidence_bytes)
    if (hashlib.sha256(body).hexdigest() != result['docx_sha256']
            or hashlib.sha256(evidence_bytes).hexdigest() != result['evidence_sha256']
            or json.loads(expected['report_manifest'].read_text(encoding='utf-8')) != result
            or evidence['project_id'] != request['project_id']
            or evidence['design_result_ref'] != result['design_result_ref']
            or evidence['design_sha256'] != source['design_sha256'] or evidence['history'] is not None):
        raise ValueError('Report evidence or checksum differs.')
    checked = ToolResult(**evidence['check'])
    _check(checked, request['project_id'], result['design_result_ref'], source['design_sha256'])
    if checked.result['summary'] != result['check_summary']:
        raise ValueError('Check summary differs.')
    with ZipFile(BytesIO(body)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 256*1024*1024 or archive.testzip() is not None:
            raise ValueError('Invalid DOCX package.')
    doc = Document(BytesIO(body))
    if (len(doc.paragraphs),len(doc.tables),len(doc.inline_shapes)) != (result['paragraphs'],result['tables'],result['images']):
        raise ValueError('DOCX contents differ from receipt.')
    return dict(path=expected['docx'], content=body, summary={key:result[key] for key in
        ('report_id','format','check_summary','paragraphs','tables','images')})
